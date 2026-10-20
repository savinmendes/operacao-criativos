# -*- coding: utf-8 -*-
"""Baixa os criativos em imagem de cada loja e monta a pasta de consulta: imagens + anuncios.md + planilha geral + folha de contato."""
import json, os, re, sys, hashlib, time, datetime as dt
import requests
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
def lim(v): return ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v
LOJAS = [("Pelvini","pelvini"),("Mockberg","mockberg"),("Venezianico","venezianico"),("AlphaGear","alphagear"),
         ("Monomods","monomods"),("Pascal Design","pascaldesign"),("Laphont","laphont"),("Marble Veil","marbleveil"),
         ("Chrono Customs","chronocustoms"),("Lucky Fours","luckyfours"),("Aston Paris","astonparis")]
FORMATOS_IMG = {"IMAGE","DCO","DPA","CAROUSEL","MULTI_IMAGES"}
HOJE = time.time()
def norm(s): return re.sub(r"[^a-z0-9]", "", (s or "").lower())
def data(ts): return dt.datetime.fromtimestamp(ts).strftime("%d/%m/%Y") if ts else ""
def dias(a):
    if not a.get("inicio"): return 0
    fim = HOJE if a.get("ativo") else (a.get("fim") or a["inicio"])
    return max(1, round((fim - a["inicio"]) / 86400))
S = requests.Session(); S.headers["User-Agent"] = "Mozilla/5.0"
wb = Workbook(); ws = wb.active; ws.title = "anuncios"
cab = ["Loja","Página","ID","Link biblioteca","Formato","Ativo","Início","Fim","Dias no ar","Variações","Plataformas","Título","Texto","Descrição","CTA","Link destino","Arquivos"]
ws.append(cab)
for c in ws[1]: c.font = Font(bold=True)
resumo = []
for nome, q in LOJAS:
    arq = f"dados/{q}.json"
    if not os.path.exists(arq): print("sem dados:", nome); continue
    ads = json.load(open(arq, encoding="utf-8"))
    da_loja = [a for a in ads if q in norm(a.get("pagina")) or norm(a.get("pagina")) in q or q in norm(__import__("urllib.parse").parse.urlparse(a.get("link") or "").netloc.split(".")[-2] if (a.get("link") and "." in __import__("urllib.parse").parse.urlparse(a.get("link")).netloc) else "")]
    pasta = os.path.join(nome.replace(" ", "-").lower()); os.makedirs(os.path.join(pasta, "imagens"), exist_ok=True)
    img_ads = [a for a in da_loja if a.get("formato") in FORMATOS_IMG and a.get("imagens")]
    img_ads.sort(key=lambda a: (-dias(a), a["id"]))
    vistos = {}; md = [f"# {nome}: criativos em imagem\n",
        f"Fonte: Biblioteca de Anúncios do Meta, busca \"{q}\", todos os países, ativos e inativos. Coleta em {dt.date.today().strftime('%d/%m/%Y')}.",
        f"Anúncios da página da loja: {len(da_loja)} (vídeo: {sum(1 for a in da_loja if a.get('formato')=='VIDEO')}; imagem/carrossel/catálogo: {len(img_ads)}).",
        "Ordem: do que ficou mais tempo no ar para o mais recente (tempo no ar é o melhor sinal de criativo que vende).\n"]
    n = 0
    for a in img_ads:
        n += 1; arqs = []
        for k, u in enumerate(a["imagens"][:10]):
            fn = f"{n:03d}_{a['formato'].lower()}_{dias(a)}d_{a['id']}" + (f"_c{k+1}" if len(a["imagens"]) > 1 else "") + ".jpg"
            cam = os.path.join(pasta, "imagens", fn)
            if os.path.exists(cam):
                h = hashlib.md5(open(cam, "rb").read()).hexdigest()
                if h in vistos: os.remove(cam); arqs.append(vistos[h] + " (repetida)")
                else: vistos[h] = fn; arqs.append(fn)
                continue
            try:
                r = S.get(u, timeout=40); r.raise_for_status()
            except Exception as e:
                arqs.append(f"(falhou: {e.__class__.__name__})"); continue
            h = hashlib.md5(r.content).hexdigest()
            if h in vistos: arqs.append(vistos[h] + " (repetida)"); continue
            try:
                im = Image.open(__import__("io").BytesIO(r.content)).convert("RGB"); im.save(os.path.join(pasta, "imagens", fn), quality=92)
            except Exception:
                open(os.path.join(pasta, "imagens", fn), "wb").write(r.content)
            vistos[h] = fn; arqs.append(fn)
        a["_arquivos"] = arqs
        md += [f"## {n:03d} · {a['formato']} · {dias(a)} dias no ar · {'ATIVO' if a.get('ativo') else 'inativo'}",
               f"- ID: [{a['id']}](https://www.facebook.com/ads/library/?id={a['id']}) · início {data(a.get('inicio'))}" + (f" · fim {data(a.get('fim'))}" if not a.get('ativo') else "") + f" · variações: {a.get('variacoes') or 1} · {', '.join(a.get('plataformas') or [])}",
               f"- Imagens: {', '.join(arqs)}",
               f"- Título: {a.get('titulo') or ''}", f"- Descrição: {a.get('link_desc') or ''}", f"- CTA: {a.get('cta') or ''} · destino: {a.get('link') or ''}",
               "- Texto:\n\n" + "\n".join("> " + l for l in (a.get('texto') or '').splitlines()) + "\n"]
        for i, c in enumerate(a.get("cards") or []):
            if c.get("titulo") or c.get("texto"):
                md.append(f"  - Card {i+1}: {c.get('titulo') or ''} | {c.get('texto') or ''}")
        md.append("")
    open(os.path.join(pasta, "anuncios.md"), "w", encoding="utf-8").write("\n".join(md))
    for a in sorted(da_loja, key=lambda a: -dias(a)):
        ws.append([lim(x) for x in [nome, a.get("pagina"), a["id"], f"https://www.facebook.com/ads/library/?id={a['id']}", a.get("formato"),
                   "sim" if a.get("ativo") else "não", data(a.get("inicio")), "" if a.get("ativo") else data(a.get("fim")), dias(a),
                   a.get("variacoes") or 1, ", ".join(a.get("plataformas") or []), a.get("titulo"), a.get("texto"), a.get("link_desc"),
                   a.get("cta"), a.get("link"), ", ".join(a.get("_arquivos", []))]])
    # folhas de contato (24 por folha) com nº e dias
    fotos = sorted(f for f in os.listdir(os.path.join(pasta, "imagens")) if f.endswith(".jpg"))
    for fo in range(0, len(fotos), 24):
        lote = fotos[fo:fo+24]; W = 300; cols = 6; rows = (len(lote)+cols-1)//cols
        folha = Image.new("RGB", (cols*W, rows*(W+28)), "white"); d = ImageDraw.Draw(folha)
        for i, f in enumerate(lote):
            im = Image.open(os.path.join(pasta, "imagens", f)); im.thumbnail((W, W))
            x, y = (i % cols)*W, (i//cols)*(W+28)
            folha.paste(im, (x + (W-im.width)//2, y + (W-im.height)//2))
            d.text((x+4, y+W+6), "_".join(f.split("_")[:3]), fill="black")
        folha.save(os.path.join(pasta, f"folha_{fo//24+1:02d}.jpg"), quality=85)
    resumo.append((nome, len(da_loja), sum(1 for a in da_loja if a.get('formato')=='VIDEO'), len(img_ads), len(fotos)))
    print(nome, "| anúncios:", len(da_loja), "| em imagem:", len(img_ads), "| imagens únicas:", len(fotos))
for col, w in zip("ABCDEFGHIJKLMNOPQ", [14,18,18,22,10,6,11,11,8,8,18,30,60,30,12,30,40]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
r = wb.create_sheet("resumo"); r.append(["Loja","Anúncios da página","Vídeo","Imagem/carrossel/catálogo","Imagens únicas baixadas"])
for x in resumo: r.append(list(x))
wb.save("anuncios_todas_lojas.xlsx")
