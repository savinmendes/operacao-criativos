# -*- coding: utf-8 -*-
"""Monta o pacote de subida no Meta (US$ 150/dia): pasta plana de imagens, plano.json, planilha, doc e moldura do catálogo."""
import json
import os
import re
import shutil

from PIL import Image, ImageDraw

from gerar_lote import GOLD, INK, CREAM, fonte, texto_c

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "SUBIDA-META-150-VENCEDORES")
IMG = os.path.join(OUT, "IMAGENS-PARA-O-GERENCIADOR")
SITE = "https://pellorini.com"
COLECAO = SITE + "/collections/best-sellers"
PDP = "PRODUTO"  # carrossel: cada card leva à página do próprio relógio
UTM = "utm_source=facebook&utm_medium=paid_social&utm_campaign={{campaign.name}}&utm_term={{adset.name}}&utm_content={{ad.name}}"

# ---------------- textos por conceito (EN) ----------------
T = {
    "oferta": {
        "textos": [
            "One for you. One for them. 🎁\nBuy any Pellorini watch and the second one is on us.\n✔ 316L stainless steel\n✔ 2-year warranty\n✔ 90-day money-back guarantee\n✔ Free tracked worldwide shipping",
            "BUY 1, GET 1 FREE ⌚⌚\nPick any two Pellorini watches. You only pay for one.\nNo code needed: the discount is applied in your cart.",
            "Why buy one watch when the second is free?\nChoose your two favorites from our best sellers and pay for just one.\nFree tracked worldwide shipping and 90 days to return.",
        ],
        "titulos": ["Buy 1 Get 1 Free", "Pick Any Two Watches", "Second Watch Free"],
        "descricoes": ["Free worldwide shipping", "No code needed"],
    },
    "preco": {
        "textos": [
            "The look people notice. The price nobody guesses. ⌚\nPellorini watches are made with 316L stainless steel and a scratch-resistant face, without the luxury markup.\nRight now, the second watch is free.",
            "Most of what you pay for a designer watch is the name on the dial.\nWe skip the logo tax: 316L steel, sapphire crystal and a 2-year warranty.\nBuy 1, get 1 free.",
            "Unpopular opinion: you don't need to spend $5,000 to look like you did.\nShop the Pellorini best sellers. Second watch free.",
        ],
        "titulos": ["The $10K Look, Without the Price Tag", "Skip the Logo Tax", "Buy 1 Get 1 Free"],
        "descricoes": ["Buy 1 Get 1 Free", "Free worldwide shipping"],
    },
    "qualidade": {
        "textos": [
            "Most cheap watches tarnish, scratch and end up in a drawer.\nPellorini is different:\n✔ 316L stainless steel that doesn't tarnish\n✔ Sapphire crystal\n✔ Adjustable bracelet with sizing tool\n✔ 2-year warranty and 90 days to return",
            "Built to be worn every day, not saved for special occasions.\n316L steel, sapphire crystal and a 2-year warranty on every watch.\nBuy 1, get 1 free.",
        ],
        "titulos": ["Built to Last. Priced to Make Sense.", "Made to Be Worn Every Day"],
        "descricoes": ["2-year warranty", "Buy 1 Get 1 Free"],
    },
    "risco": {
        "textos": [
            "Wear it for 90 days.\nIf you don't love it, send it back for a full refund. No questions asked.\nEvery Pellorini also comes with a 2-year warranty and free tracked worldwide shipping.",
            "Buying a watch online feels like a risk. So we removed it.\n90-day money-back guarantee + 2-year warranty.\nAnd today, the second watch is free.",
        ],
        "titulos": ["90-Day Money-Back Guarantee", "Love It or Return It"],
        "descricoes": ["Buy 1 Get 1 Free", "2-year warranty"],
    },
    "desejo_h": {
        "textos": [
            "Meet our best sellers. ⌚\nSunray dials, 316L stainless steel and a bracelet that fits any wrist.\nPick any two, the second one is free.",
            "Dressed up or down, it just works.\nDiscover the Pellorini best sellers with free tracked worldwide shipping.",
            "The watch people ask about.\nShop Pellorini: buy 1 and get 1 free.",
        ],
        "titulos": ["Meet the Pellorini Best Sellers", "Your New Everyday Watch", "Buy 1 Get 1 Free"],
        "descricoes": ["Buy 1 Get 1 Free", "Free worldwide shipping"],
    },
    "desejo_m": {
        "textos": [
            "Small details. Big difference. ✨\nDelicate gold-tone watches made for every day.\nPick any two, the second one is free.",
            "The everyday piece she'll never take off.\nFree tracked worldwide shipping and 90 days to return.",
            "Quiet luxury, every day.\nShop our women's best sellers: buy 1, get 1 free.",
        ],
        "titulos": ["Everyday Elegance", "Buy 1 Get 1 Free", "The Gift She'll Wear Every Day"],
        "descricoes": ["Free worldwide shipping", "Buy 1 Get 1 Free"],
    },
    "presente": {
        "textos": [
            "The gift that comes in pairs. 🎁\nBuy a Pellorini for them and the second one is free. Keep it or gift it.\nFree tracked worldwide shipping.",
            "Anniversary, birthday or just because.\nBuy 1 Pellorini watch, get 1 free: one for you, one for them.",
        ],
        "titulos": ["One for You. One for Them.", "Buy 1 Get 1 Free"],
        "descricoes": ["Free worldwide shipping", "90-day returns"],
    },
    "escolha": {
        "textos": [
            "Classic or iced? Silver or gold?\nYou don't have to choose: buy 1 Pellorini, get the second one free.",
            "Which one are you?\nPick your two favorites from our best sellers and pay for just one.",
        ],
        "titulos": ["Can't Choose? Get Both.", "Buy 1 Get 1 Free"],
        "descricoes": ["Free worldwide shipping", "No code needed"],
    },
}

L1, L2 = "lote1", "lote2"


def imgs(pasta, *nomes):
    return [(pasta, n) for n in nomes]


# ---------------- estrutura ----------------
L3 = "lote3-vencedores"
W5 = ["Vittorio", "Dominante", "Milano", "Opulent", "Maretti"]  # ranking de vendas da loja

def por_modelo(fam, prefixo, nomes=W5):
    return imgs(f"{L3}/{fam}", *[f"{prefixo}_{n}" for n in nomes])

CONJUNTOS = [
    ("CJ1", "Masculino | Oferta", [
        ("A01", "C1_P2_bogo-saturado", "oferta", COLECAO, imgs(L3 + "/bogo-saturado", "C1_P2_bogo-saturado_Vittorio-Milano", "C1_P2_bogo-saturado_Dominante-Opulent", "C1_P2_bogo-saturado_Milano-Maretti", "C1_P2_bogo-saturado_Maretti-Dominante")),
        ("A02", "C2_P1_vs-designer", "preco", PDP, por_modelo("vs-designer", "C2_P1_vs-designer")),
        ("A03", "C2_P1_is-name-worth", "preco", PDP, por_modelo("is-name-worth", "C2_P1_is-name-worth")),
        ("A04", "C8_P2_postit", "oferta", PDP, por_modelo("postit", "C8_P2_postit-produto") + imgs(L3 + "/postit", "C8_P2_postit_pelvini-milano_5")),
    ]),
    ("CJ2", "Masculino | Desejo", [
        ("A05", "C7_P1_poster", "desejo_h", PDP, por_modelo("poster", "C7_P1_poster")),
        ("A06", "C2_P1_pulso-frase", "desejo_h", PDP, imgs(L3 + "/pulso-frase", "C2_P1_pulso-frase_vittorio_5", "C2_P1_pulso-frase_pelvini-milano_5") + imgs(L3 + "/vida-bogo", "C1_P2_vida-bogo_vittorio_5", "C1_P2_vida-bogo_pelvini-milano_5")),
        ("A07", "C7_P1_split", "desejo_h", PDP, imgs(L3 + "/split", "C7_P1_split_vittorio_5", "C7_P1_split_pelvini-milano_5")),
        ("A08", "C7_P1_caixa-maos", "desejo_h", PDP, imgs(L3 + "/caixa-maos", "C7_P1_caixa-maos_vittorio_2", "C7_P1_caixa-maos_pelvini-dominante_2", "C7_P1_caixa-maos_pelvini-milano_2", "C7_P1_caixa-maos_pelvini-opulent_2", "C7_P1_caixa-maos_pelvini-maretti-1_2")),
    ]),
    ("CJ3", "Presente", [
        ("A09", "C1_P2_imessage", "presente", COLECAO, imgs(L3 + "/imessage", "C1_P2_imessage_pai-dominante", "C1_P2_imessage_irmao-milano")),
        ("A10", "C1_P2_his-hers", "presente", COLECAO, imgs(L3 + "/his-hers", "C1_P2_his-hers_vittorio-clovera", "C1_P2_his-hers_milano-serafina", "C1_P2_his-hers_opulent-lunetta")),
        ("A11", "C8_P2_maleta", "oferta", COLECAO, imgs(L3 + "/maleta", "C8_P2_maleta_bestsellers") + imgs(L2 + "/maleta", "C8_P5_maleta_mulher")),
        ("A12", "C7_P1_qual-voce", "escolha", COLECAO, imgs(L3 + "/qual-voce", "C7_P1_qual-voce_classic-bold", "C7_P1_qual-voce_green-blue", "C7_P1_qual-voce_chrono-clean", "C7_P1_qual-voce_grey-green")),
    ]),
    ("CJ4", "Feminino", [
        ("A13", "C7_P5_poster", "desejo_m", PDP, imgs(L2 + "/poster", "C7_P5_poster_Clovera", "C7_P5_poster_Serafina", "C7_P5_poster_Lunetta", "C7_P5_poster_Fiorella", "C7_P5_poster_Grazia")),
        ("A14", "C7_P5_pulso-frase", "desejo_m", PDP, imgs(L2 + "/pulso-frase", "C7_P5_pulso-frase_pelvini-clovera_3", "C7_P5_pulso-frase_pelvini-aria_4", "C7_P5_pulso-frase_pelvini-bellissima_6", "C7_P5_pulso-frase_pelvini-carina_4")),
        ("A15", "C7_P5_split", "desejo_m", PDP, imgs(L2 + "/split", "C7_P5_split_pelvini-clovera_4", "C7_P5_split_pelvini-aria_3", "C7_P5_split_pelvini-bellissima_3", "C7_P5_split_pelvini-carina_6")),
        ("A16", "C1_P5_vida-bogo", "oferta", PDP, imgs(L2 + "/vida-bogo", "C1_P2_vida-bogo_pelvini-aria_4", "C1_P2_vida-bogo_pelvini-clovera_3", "C1_P2_vida-bogo_pelvini-carina_3", "C1_P2_vida-bogo_pelvini-bellissima_4")),
    ]),
    ("CJ5", "Confiança", [
        ("A17", "C4_P3_garantia-90d", "risco", PDP, por_modelo("garantia-90d", "C4_P3_garantia-90d")),
        ("A18", "C3_P3_generico-x-nos", "qualidade", PDP, por_modelo("generico-x-nos", "C3_P3_generico-x-nos")),
        ("A19", "C2_P1_tweet", "preco", PDP, por_modelo("tweet", "C2_P1_tweet")),
        ("A20", "C2_P1_ancora-10k", "preco", PDP, por_modelo("ancora-10k", "C2_P1_ancora-10k")),
    ]),
]

CAMP_VENDAS = {
    "nome": "PELLORINI | Vendas | BOGO | EN | ABO",
    "objetivo": "Vendas (OUTCOME_SALES)",
    "tipo_orcamento": "ABO (orçamento no conjunto): garante verba a cada conceito no teste",
    "status_inicial": "PAUSADA",
    "categoria_especial": "Nenhuma",
}
CONJ_PADRAO = {
    "orcamento_diario_usd": 20,
    "otimizacao": "Conversões > Compra (Purchase) no site",
    "lance": "Maior volume, sem limite de custo nos dias 1–14. A partir do dia 15: limite de custo de US$ 20",
    "atribuicao": "7 dias após o clique ou 1 dia após a visualização",
    "paises": ["US", "GB", "CA", "AU"],
    "idade": "18–65+",
    "genero": "Todos",
    "publico": "Advantage+ audience ligado, sem interesses (o criativo segmenta)",
    "posicionamentos": "Advantage+ (todos)",
    "idioma": "Sem filtro",
}
CAMP_CATALOGO = {
    "nome": "PELLORINI | Catálogo | BOGO | EN",
    "objetivo": "Vendas (OUTCOME_SALES) com anúncio de catálogo (Advantage+ catalog ads)",
    "orcamento_diario_usd": 50,
    "status_inicial": "PAUSADA",
    "catalogo": "Catálogo da loja Pellorini na Shopify (pellorini.com / essiora.com) > conjunto de produtos com a tag 'best-sellers' (45 produtos)",
    "conjunto": "CJ6 | Catálogo | Best Sellers: mesmo público, países, otimização e lance do padrão; orçamento US$ 50/dia",
    "anuncio": "A21 | C9_catalogo-bogo: anúncio de catálogo (carrossel + coleção) com a moldura 'moldura-catalogo-bogo.png' aplicada às imagens do catálogo",
    "textos": T["oferta"],
}
REGRAS = [
    "Tudo criado PAUSADO. Só ativar depois da aprovação da pessoa.",
    "Melhorias do Advantage+ Creative: desligar reescrita/variação de texto e efeitos automáticos; manter ajuste de posicionamento.",
    "Anúncios com formato 'flexivel' (vários relógios por imagem): asset feed com todas as imagens (1:1 e 9:16) e todas as variações de texto, título e descrição; link = coleção.",
    "Anúncios com formato 'carrossel' (1 relógio por imagem): cada card leva à página do próprio relógio (cards[].link), com título = nome do relógio e descrição 'Buy 1 Get 1 Free'; texto principal = textos_principais[0].",
    "Cortar um anúncio que gastar US$ 40 (2× o CPA) sem compra. Julgar o conjunto a cada 7 dias.",
    "Escalar +20% a cada ~4 dias nos conjuntos com CPA de 3–7 dias ≤ US$ 20; se passar, −15%.",
    "Família nova entra em conjunto novo; não pausar vencedores ao subir novos.",
]


NOMES_H = {"Vittorio": "vittorio", "Dominante": "pelvini-dominante", "Milano": "pelvini-milano", "Opulent": "pelvini-opulent",
           "Maretti": "pelvini-maretti-1", "Clovera": "pelvini-clovera", "Serafina": "pelvini-serafina", "Lunetta": "pelvini-lunetta",
           "Fiorella": "pelvini-fiorella", "Grazia": "pelvini-grazia", "Titan": "titan", "Lucent": "pelvini-lucent", "Dante": "pelvini-dante"}


def handle_de(base):
    m = re.search(r"_(pelvini-[a-z0-9-]+?|vittorio|titan)_\d+$", base)
    if m:
        return m.group(1)
    return NOMES_H[base.rsplit("_", 1)[-1]]


MAPA_SLUG = json.load(open(os.path.join(BASE, "..", "PELLORINI-LOJA", "slugs", "mapa_slugs.json")))


def titulo_de(handle):
    return "Pellorini " + handle.replace("pelvini-", "").replace("-1", "").replace("-", " ").title()


def moldura():
    w = h = 1080
    im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((40, 40, 380, 150), radius=16, fill=GOLD + (255,))
    d.text((66, 52), "BUY 1", font=fonte("anton", 44), fill=INK)
    d.text((66, 98), "GET 1 FREE", font=fonte("anton", 40), fill=INK)
    d.rectangle((0, h - 90, w, h), fill=(0, 0, 0, 235))
    texto_c(d, w, h - 62, "FREE WORLDWIDE SHIPPING  ·  90-DAY RETURNS", fonte("sans", 26, 600), CREAM, espaco=2)
    return im


def main():
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(IMG)
    plano = {"site": SITE, "parametros_url": UTM, "campanha_vendas": CAMP_VENDAS, "padrao_conjunto": CONJ_PADRAO,
             "conjuntos": [], "campanha_catalogo": CAMP_CATALOGO, "regras": REGRAS, "textos_por_conceito": T,
             "cta": "SHOP_NOW (Comprar agora)"}
    linhas = []
    falta = []
    for cid, cnome, ads in CONJUNTOS:
        cj = {"id": cid, "nome": f"{cid} | {cnome}", **CONJ_PADRAO, "anuncios": []}
        for aid, anome, conceito, url, lista in ads:
            arquivos, cards = [], []
            fmts = ("1x1",) if url == PDP else ("1x1", "9x16")
            for pasta, base in lista:
                for fmt in fmts:
                    src = os.path.join(BASE, pasta, f"{base}_{fmt}.jpg")
                    if not os.path.exists(src):
                        falta.append(src)
                        continue
                    dst = f"{aid}__{base}_{fmt}.jpg".replace(" ", "-")
                    shutil.copy2(src, os.path.join(IMG, dst))
                    arquivos.append(dst)
                    if url == PDP:
                        h = handle_de(base)
                        cards.append({"imagem": dst, "link": f"{SITE}/products/{MAPA_SLUG.get(h, h)}", "titulo": titulo_de(h), "descricao": "Buy 1 Get 1 Free"})
            t = T[conceito]
            ad = {"id": aid, "nome": f"{aid} | {anome}", "conceito": conceito, "cta": "SHOP_NOW",
                  "textos_principais": t["textos"], "titulos": t["titulos"], "descricoes": t["descricoes"]}
            if url == PDP:
                ad.update({"formato": "carrossel", "url": COLECAO, "cards": cards, "imagens": arquivos,
                           "obs": "Carrossel: cada card com a própria imagem 1:1, título, descrição e link da página do relógio. Texto principal = textos_principais[0] (carrossel aceita 1). Link do card final 'ver mais' = url (coleção). Desligar 'ordenar cards automaticamente' não é necessário."})
                destino = "Carrossel: cada card → página do produto" + chr(10) + chr(10).join(f"{c['titulo']}: {c['link']}" for c in cards)
            else:
                ad.update({"formato": "flexivel", "url": url, "imagens": arquivos})
                destino = url
            cj["anuncios"].append(ad)
            linhas.append([CAMP_VENDAS["nome"], cj["nome"], f"{aid} | {anome}", conceito, destino, len(arquivos), chr(10).join(arquivos),
                           (chr(10) + "---" + chr(10)).join(t["textos"]), " / ".join(t["titulos"]), " / ".join(t["descricoes"]), "SHOP_NOW"])
        plano["conjuntos"].append(cj)
    moldura().save(os.path.join(IMG, "A21__moldura-catalogo-bogo.png"))
    linhas.append([CAMP_CATALOGO["nome"], "CJ6 | Catálogo | Best Sellers", "A21 | C9_catalogo-bogo", "oferta", "Página de cada produto (catálogo)", 1,
                   "A21__moldura-catalogo-bogo.png", "\n---\n".join(T["oferta"]["textos"]), " / ".join(T["oferta"]["titulos"]), " / ".join(T["oferta"]["descricoes"]), "SHOP_NOW"])
    json.dump(plano, open(os.path.join(OUT, "plano-subida.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "anuncios"
    ws.append(["Campanha", "Conjunto", "Anúncio", "Conceito", "Link de destino", "Nº imagens", "Arquivos", "Textos principais", "Títulos", "Descrições", "CTA"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for l in linhas:
        ws.append(l)
    for col, wd in zip("ABCDEFGHIJK", [30, 26, 30, 12, 42, 8, 50, 70, 40, 30, 10]):
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    w2 = wb.create_sheet("configuracoes")
    for k, v in list(CAMP_VENDAS.items()) + [("—", "—")] + list(CONJ_PADRAO.items()) + [("—", "—")] + list(CAMP_CATALOGO.items()):
        w2.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v])
    w2.append(["parametros_url", UTM])
    for r in REGRAS:
        w2.append(["regra", r])
    w2.column_dimensions["A"].width = 26
    w2.column_dimensions["B"].width = 120
    wb.save(os.path.join(OUT, "PLANILHA-SUBIDA.xlsx"))
    print("imagens:", len(os.listdir(IMG)), "| anúncios:", len(linhas), "| faltando:", len(falta))
    for f in falta:
        print("  falta:", f)


if __name__ == "__main__":
    main()
